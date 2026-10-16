#!/usr/bin/env python3
"""Sauvegarde de l'espace abonnés : fabrique les deux fichiers à garder hors du dépôt.

  Sauvegarde_espace_abonnes_<date>.xlsx  version lisible (abonnés, abonnements, suspensions, contenus)
  Sauvegarde_espace_abonnes_<date>.sql   remise des données dans la base après une panne

Mode d'emploi
  1. Lancer la requête de tools/sauvegarde.sql avec le connecteur Supabase (lecture seule).
  2. Enregistrer le résultat dans un fichier (le connecteur le fait seul quand il est long).
  3. python3 -I tools/sauvegarde.py <fichier du résultat> <dossier de sortie> [--drive]

Avec --drive, le script prépare aussi un sous-dossier « drive » de fichiers texte courts, à envoyer un par un
avec le connecteur Google Drive (qui ne transmet que du texte recopié) : un CSV par tableau, converti en feuille
Google Sheets à l'envoi, et le fichier de restauration découpé en parties autonomes, à lancer dans l'ordre.

Le dossier de sortie doit être HORS du dépôt (il est public) : ces fichiers contiennent les noms,
téléphones et e-mails des abonnés. Nécessite openpyxl.
"""
import csv
import hashlib
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre"]


def lire_resultat(chemin):
    """Accepte le JSON brut de la requête ou la réponse complète du connecteur (texte encadré)."""
    brut = Path(chemin).read_text(encoding="utf8")
    try:
        brut = json.loads(brut)
        if isinstance(brut, dict) and "result" in brut:
            brut = brut["result"]
    except json.JSONDecodeError:
        pass
    if isinstance(brut, str):
        # Le connecteur entoure les lignes d'une balise <untrusted-data-…> citée une première fois dans
        # sa phrase d'avertissement : les lignes suivent la balise placée en début de ligne.
        debut = brut.find("\n<untrusted-data-")
        if debut >= 0:
            brut = brut[brut.index(">", debut) + 1:]
        brut, _ = json.JSONDecoder().raw_decode(brut.lstrip())
    if isinstance(brut, list):
        brut = brut[0]
    if "sauvegarde" in brut:
        brut = brut["sauvegarde"]
    if isinstance(brut, str):
        brut = json.loads(brut)
    return {k: brut.get(k) or [] for k in ("utilisateurs", "profils", "admins", "abonnements",
                                          "suspensions", "contenus")}


def heure(s):
    """Horodatage en heure d'Abidjan (GMT)."""
    if not s:
        return ""
    return datetime.fromisoformat(s).astimezone(timezone.utc).strftime("%d/%m/%Y %H:%M")


def date(s):
    return datetime.fromisoformat(s).strftime("%d/%m/%Y") if s else ""


MARINE = PatternFill("solid", fgColor="0F2A5C")
ENTETE = Font(bold=True, color="FFFFFF")


def feuille(wb, titre, entetes, lignes, largeurs):
    ws = wb.create_sheet(titre)
    ws.append(entetes)
    for c in ws[1]:
        c.fill, c.font = MARINE, ENTETE
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for ligne in lignes:
        ws.append(ligne)
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def tableaux(d, jour):
    """Les tableaux lisibles, communs au classeur Excel et aux CSV : (titre, en-têtes, lignes, largeurs)."""
    emails = {u["id"]: u["email"] for u in d["utilisateurs"]}
    profils = {p["id"]: p for p in d["profils"]}
    admins = {a["user_id"] for a in d["admins"]}
    nom_de = lambda uid: f"{profils.get(uid, {}).get('prenom', '')} {profils.get(uid, {}).get('nom', '')}".strip()

    dernier = {}
    for b in d["abonnements"]:
        if not b["supprime_le"] and (b["user_id"] not in dernier or b["fin"] > dernier[b["user_id"]]["fin"]):
            dernier[b["user_id"]] = b
    abonnes = []
    for u in d["utilisateurs"]:
        p = profils.get(u["id"], {})
        b = dernier.get(u["id"])
        if u["id"] in admins:
            etat = "Administrateur"
        elif b:
            etat = "En cours" if b["debut"] <= jour <= b["fin"] else ("À venir" if jour < b["debut"] else "Terminé")
        else:
            etat = "Sans abonnement"
        abonnes.append([p.get("prenom", ""), p.get("nom", ""), u["email"], p.get("telephone", ""),
                        p.get("concours") or "", heure(u["cree_le"]), heure(u["derniere_connexion"]),
                        date(b["fin"]) if b else "", etat, u["id"]])

    return [
        ("Abonnés", ["Prénom", "Nom", "E-mail", "Téléphone", "Concours", "Inscrit le", "Dernière connexion",
                     "Accès jusqu'au", "État", "Identifiant (technique)"],
         abonnes, [16, 18, 30, 18, 16, 18, 18, 15, 16, 38]),
        ("Abonnements", ["N°", "Abonné", "E-mail", "Début", "Fin (incluse)", "Montant (F CFA)", "Moyen", "Note",
                         "Activé le", "Activé par", "Statut"],
         [[b["id"], nom_de(b["user_id"]), emails.get(b["user_id"], ""), date(b["debut"]), date(b["fin"]),
           b["montant"], b["moyen"] or "", b["note"] or "", heure(b["cree_le"]), emails.get(b["cree_par"], ""),
           "Supprimé le " + heure(b["supprime_le"]) if b["supprime_le"] else "Valable"] for b in d["abonnements"]],
         [6, 24, 30, 12, 13, 15, 14, 24, 18, 24, 26]),
        ("Suspensions", ["N°", "Abonné", "E-mail", "Début", "Réactivé le", "Créée le", "Créée par"],
         [[s["id"], nom_de(s["user_id"]), emails.get(s["user_id"], ""), date(s["debut"]), date(s["fin"]),
           heure(s["cree_le"]), emails.get(s["cree_par"], "")] for s in d["suspensions"]],
         [6, 24, 30, 12, 12, 18, 24]),
        # Une cellule Excel ne peut dépasser 32 767 caractères : le texte complet reste dans le fichier .sql.
        ("Contenus", ["N°", "Titre", "Matière", "Type", "Concours", "Publié", "Publié le",
                      "Fichier (compartiment exclusif)", "Lien", "Description", "Texte (QCM, fiches)"],
         [[c["id"], c["titre"], c["matiere"] or "", c["type"], ", ".join(c["concours"] or []),
           "Oui" if c["publie"] else "Non", heure(c["publie_le"]), c["fichier"] or "", c["lien"] or "",
           c["description"] or "", (c["texte"] or "")[:32000]] for c in d["contenus"]],
         [6, 50, 28, 8, 16, 8, 18, 50, 20, 50, 40]),
    ]


def classeur(d, jour, jour_long, nom):
    fichiers = sum(1 for c in d["contenus"] if c["fichier"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Lisez-moi"
    texte = [
        ("Sauvegarde de l'espace abonnés coursesesame.com", True),
        (f"Faite le {jour_long}, depuis la base Supabase sesame-saj-espace-abonnes.", False),
        ("", False),
        ("Contenu", True),
        (f"Abonnés. {len(d['profils'])} comptes, avec e-mail, téléphone et concours.", False),
        (f"Abonnements. {len(d['abonnements'])} périodes payées (montants, moyens de paiement, dates).", False),
        (f"Suspensions. {len(d['suspensions'])}.", False),
        (f"Contenus. {len(d['contenus'])} contenus exclusifs (titres, matières, chemins des fichiers, "
         "textes des QCM).", False),
        ("", False),
        ("Ce que ce fichier ne contient pas", True),
        (f"Les PDF eux-mêmes ({fichiers} fichiers du compartiment exclusif). Gardez vos PDF d'origine à part.",
         False),
        ("Les mots de passe des abonnés. Supabase les garde chiffrés, ils ne sont pas exportables.", False),
        ("", False),
        ("Confidentialité", True),
        ("Ce fichier contient des données personnelles d'abonnés. Ne le partagez avec personne et ne "
         "l'envoyez pas sur WhatsApp.", False),
        ("", False),
        ("Restauration", True),
        (f"Le fichier {nom}.sql permet de remettre ces données dans la base en cas de panne. "
         "À faire avec Claude, de préférence.", False),
        ("Les dates de début et de fin d'abonnement sont des dates d'Abidjan, incluses. Les autres heures "
         "sont en heure GMT (heure d'Abidjan).", False),
    ]
    for t, gras in texte:
        ws.append([t])
        if gras:
            ws.cell(ws.max_row, 1).font = Font(bold=True, color="0F2A5C", size=13 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 120

    for titre, entetes, lignes, largeurs in tableaux(d, jour):
        ws = feuille(wb, titre, entetes, lignes, largeurs)
        if titre == "Abonnements":
            ws.append([])
            ws.append(["", "Total des abonnements valables", "", "", "",
                       sum(b["montant"] or 0 for b in d["abonnements"] if not b["supprime_le"])])
            ws.cell(ws.max_row, 2).font = ws.cell(ws.max_row, 6).font = Font(bold=True)
            for (c,) in ws.iter_rows(min_row=2, min_col=6, max_col=6):
                c.number_format = "# ##0"
    return wb


def bloc(table, lignes, identite):
    if not lignes:
        return f"-- {table} : aucune ligne.\n\n"
    js = json.dumps(lignes, ensure_ascii=False)
    assert "$sauvegarde$" not in js
    s = (f"insert into public.{table}{' overriding system value' if identite else ''}\n"
         f"select * from json_populate_recordset(null::public.{table}, $sauvegarde${js}$sauvegarde$)\n"
         "on conflict do nothing;\n")
    if identite:
        s += (f"select setval(pg_get_serial_sequence('public.{table}', 'id'), "
              f"greatest((select max(id) from public.{table}), 1));\n")
    return s + "\n"


def entete(jour_long):
    return (f"-- Sauvegarde de l'espace abonnés coursesesame.com, faite le {jour_long}.\n"
           "-- Projet Supabase sesame-saj-espace-abonnes (iyftcjszvnaaakirlsdd).\n"
           "-- Contient des données personnelles : ne pas partager.\n"
           "--\n"
           "-- Usage : remettre ces lignes dans la base après une panne, à faire de préférence avec Claude.\n"
           "-- Conditions : la structure de la base doit exister (dossier supabase/migrations du dépôt), et les\n"
           "-- comptes doivent avoir été recréés dans Authentication avec les mêmes identifiants (liste\n"
           "-- ci-dessous), sinon les profils ne peuvent pas leur être rattachés.\n"
           "-- Les lignes déjà présentes sont ignorées (on conflict do nothing) : on peut relancer sans doublon.\n"
           "--\n")


def comptes(d):
    return ("-- Comptes au moment de la sauvegarde (identifiant, e-mail) :\n"
            + "".join(f"--   {u['id']}  {u['email']}\n" for u in d["utilisateurs"]))


def restauration(d, jour_long):
    sql = entete(jour_long) + comptes(d) + "\nbegin;\n\n"
    sql += bloc("profils", d["profils"], False)
    sql += bloc("admins", d["admins"], False)
    sql += bloc("abonnements", d["abonnements"], True)
    sql += bloc("suspensions", d["suspensions"], True)
    sql += bloc("contenus", d["contenus"], True)
    return sql + "commit;\n"


def restauration_parties(d, jour_long, taille=30000):
    """Le fichier de restauration en parties autonomes de taille limitée (chacune dans sa propre transaction).
    La première porte les comptes, les abonnements et les suspensions ; les suivantes, les contenus."""
    groupes, groupe, poids = [], [], 0
    for c in d["contenus"]:
        n = len(json.dumps(c, ensure_ascii=False))
        if groupe and poids + n > taille:
            groupes.append(groupe)
            groupe, poids = [], 0
        groupe.append(c)
        poids += n
    if groupe:
        groupes.append(groupe)
    total = 1 + len(groupes)
    titre = lambda k, objet: (entete(jour_long) + f"-- Partie {k} sur {total} ({objet}). Lancer les parties "
                              "dans l'ordre ; chacune est autonome.\n--\n")
    parties = [titre(1, "comptes, abonnements, suspensions") + comptes(d) + "\nbegin;\n\n"
               + bloc("profils", d["profils"], False) + bloc("admins", d["admins"], False)
               + bloc("abonnements", d["abonnements"], True) + bloc("suspensions", d["suspensions"], True)
               + "commit;\n"]
    for k, g in enumerate(groupes, 2):
        parties.append(titre(k, f"contenus {g[0]['id']} à {g[-1]['id']}") + "\nbegin;\n\n"
                       + bloc("contenus", g, True) + "commit;\n")
    return parties


def pour_drive(d, jour, jour_long, nom, dossier):
    """Fichiers texte courts pour le connecteur Google Drive. Renvoie (nom, taille, empreinte) de chacun."""
    dossier.mkdir(parents=True, exist_ok=True)
    for vieux in dossier.glob(f"{nom}_*"):
        vieux.unlink()
    fichiers = []
    for titre, entetes, lignes, _ in tableaux(d, jour):
        if not lignes:
            continue
        if titre == "Contenus":  # les textes complets, trop longs pour un tableau, sont dans la restauration
            entetes, lignes = entetes[:-1], [l[:-1] for l in lignes]
        tampon = io.StringIO()
        csv.writer(tampon, lineterminator="\n").writerows([entetes] + lignes)
        fichiers.append((f"{nom}_{titre.lower().replace('é', 'e')}.csv", tampon.getvalue()))
    parties = restauration_parties(d, jour_long)
    for k, texte in enumerate(parties, 1):
        fichiers.append((f"{nom}_restauration_{k}_sur_{len(parties)}.sql", texte))
    resume = []
    for f, texte in fichiers:
        octets = texte.encode("utf8")
        (dossier / f).write_bytes(octets)
        resume.append((f, len(octets), hashlib.sha256(octets).hexdigest()))
    return resume


def main():
    drive = "--drive" in sys.argv[1:]
    args = [a for a in sys.argv[1:] if a != "--drive"]
    if len(args) != 2:
        sys.exit("Usage : python3 -I tools/sauvegarde.py <fichier du résultat> <dossier de sortie> [--drive]")
    d = lire_resultat(args[0])
    sortie = Path(args[1]).resolve()
    depot = Path(__file__).resolve().parent.parent
    if sortie == depot or depot in sortie.parents:
        sys.exit("Le dossier de sortie doit être hors du dépôt : la sauvegarde contient des données personnelles.")
    sortie.mkdir(parents=True, exist_ok=True)

    maintenant = datetime.now(timezone.utc)  # Abidjan vit à l'heure GMT toute l'année
    jour = maintenant.strftime("%Y-%m-%d")
    jour_long = f"{maintenant.day} {MOIS[maintenant.month - 1]} {maintenant.year}"
    nom = f"Sauvegarde_espace_abonnes_{jour}"

    classeur(d, jour, jour_long, nom).save(sortie / f"{nom}.xlsx")
    (sortie / f"{nom}.sql").write_text(restauration(d, jour_long), encoding="utf8")

    valables = [b for b in d["abonnements"] if not b["supprime_le"]]
    en_cours = sum(1 for b in valables if b["debut"] <= jour <= b["fin"])
    print(f"{sortie / nom}.xlsx")
    print(f"{sortie / nom}.sql")
    if drive:
        for f, taille, empreinte in pour_drive(d, jour, jour_long, nom, sortie / "drive"):
            print(f"{sortie / 'drive' / f}  {taille} octets  sha256 {empreinte}")
    print(f"{len(d['utilisateurs'])} comptes, {len(d['admins'])} administrateur(s), "
          f"{len(d['abonnements'])} abonnements dont {en_cours} en cours, "
          f"{sum(b['montant'] or 0 for b in valables)} F au total, {len(d['suspensions'])} suspension(s), "
          f"{len(d['contenus'])} contenus dont {sum(1 for c in d['contenus'] if c['fichier'])} fichiers.")


if __name__ == "__main__":
    main()
